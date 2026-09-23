import pytest
from pathlib import Path

from rag.models import IngestError
from rag.parse import parse_file


def test_markdown_offsets_heading_code_and_table(tmp_path):
    path = tmp_path / "guide.md"
    path.write_text(
        "# Dosage\n\nTake 5 mg.\n\n"
        "| sku | name |\n| --- | --- |\n| SKU-1 | housing |\n\n"
        "```\nexport PATH=/opt/bin\n```\n",
        encoding="utf-8",
    )
    mime, blocks = parse_file(path)
    assert mime == "text/markdown"
    document = "\n".join(block.text for block in blocks)
    for block in blocks:
        assert document[block.start_char : block.end_char] == block.text
    assert blocks[0].heading_path == "Dosage"
    assert "5 mg" in blocks[0].text
    assert any(block.kind == "table" and "SKU-1" in block.text for block in blocks)
    assert any(block.kind == "code" and "PATH" in block.text for block in blocks)


def test_txt_uses_the_filename_as_the_heading(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("First paragraph.\n\nSecond paragraph.\n", encoding="utf-8")
    _mime, blocks = parse_file(path)
    assert [block.heading_path for block in blocks] == ["notes", "notes"]
    assert blocks[0].text == "First paragraph."


def test_html_drops_script_and_keeps_the_heading(tmp_path):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><script>secret()</script><h1>Dosage</h1><p>Take 5 mg.</p></html>",
        encoding="utf-8",
    )
    _mime, blocks = parse_file(path)
    assert "secret" not in " ".join(block.text for block in blocks)
    assert blocks[0].heading_path == "Dosage"
    assert blocks[0].text == "Take 5 mg."


def test_html_div_only_page_keeps_text(tmp_path):
    path = tmp_path / "modern.html"
    path.write_text(
        "<html><body><div><h4>Scope</h4><section>Applies to all campuses.</section>"
        '<img alt="Campus map"></div></body></html>',
        encoding="utf-8",
    )
    _mime, blocks = parse_file(path)
    assert any(block.heading_path == "Scope" for block in blocks)
    assert any("Applies to all campuses." in block.text for block in blocks)
    assert any(block.kind == "caption" and "Campus map" in block.text for block in blocks)


def test_html_with_only_a_script_is_refused(tmp_path):
    path = tmp_path / "empty.html"
    path.write_text("<html><script>secret()</script></html>", encoding="utf-8")
    with pytest.raises(IngestError, match="no text"):
        parse_file(path)


def test_csv_repeats_the_column_name(tmp_path):
    path = tmp_path / "parts.csv"
    path.write_text("sku,name\nSKU-7842-XL,housing\n", encoding="utf-8")
    _mime, blocks = parse_file(path)
    assert blocks[0].kind == "row"
    assert "sku: SKU-7842-XL" in blocks[0].text
    assert blocks[0].heading_path == "parts"


def test_pdf_round_trip(tmp_path):
    path = tmp_path / "note.pdf"
    path.write_bytes(_pdf("Hello PDF"))
    _mime, blocks = parse_file(path)
    assert any("Hello PDF" in block.text for block in blocks)
    assert blocks[0].page == 1


def test_pdf_tables_are_not_also_kept_as_prose():
    # F-02: distinctive table figures must not also sit in prose on the same page.
    path = Path("documents/Carbon_Reduction_Plan.pdf")
    if not path.exists():
        pytest.skip("corpus PDF missing")
    _mime, blocks = parse_file(path)
    markers = ("14,644", "914.38", "8,688", "5,543", "37,445", "852.95")
    for marker in markers:
        homes = [(block.page, block.kind) for block in blocks if marker in block.text]
        assert homes, f"missing {marker}"
        kinds = {kind for _page, kind in homes}
        assert "prose" not in kinds, f"{marker} still indexed as prose: {homes}"
        assert "table" in kinds, f"{marker} not in a table block: {homes}"


def test_pdf_without_a_text_layer_is_refused(tmp_path):
    path = tmp_path / "blank.pdf"
    path.write_bytes(_pdf(""))
    with pytest.raises(IngestError, match="no text"):
        parse_file(path)


def test_docx_heading_and_table(tmp_path):
    from docx import Document

    path = tmp_path / "note.docx"
    document = Document()
    document.add_heading("Dosage", level=1)
    document.add_paragraph("Take 5 mg.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "sku"
    table.rows[0].cells[1].text = "name"
    table.rows[1].cells[0].text = "SKU-1"
    table.rows[1].cells[1].text = "housing"
    document.save(path)
    _mime, blocks = parse_file(path)
    assert any(block.heading_path == "Dosage" and "5 mg" in block.text for block in blocks)
    assert any(block.kind == "table" and "SKU-1" in block.text for block in blocks)


def _pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    if escaped:
        content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET\n".encode("latin-1")
    else:
        content = b"BT ET\n"
    stream = b"<< /Length %d >>\nstream\n" % len(content) + content + b"endstream"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        stream,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 6\n0000000000 65535 f \n"
    for number in range(1, 6):
        out += f"{offsets[number]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)
