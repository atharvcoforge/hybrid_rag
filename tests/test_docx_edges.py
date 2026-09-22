"""DOCX edge cases: footnotes, text boxes, tracked changes (§10)."""

import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from rag.parse import parse_file


def test_docx_keeps_list_items(tmp_path):
    path = tmp_path / "lists.docx"
    document = Document()
    document.add_heading("Duties", level=1)
    document.add_paragraph("First duty", style="List Bullet")
    document.add_paragraph("Second duty", style="List Number")
    document.save(path)
    _mime, blocks = parse_file(path)
    bodies = [block.text for block in blocks if block.kind == "list"]
    assert "First duty" in bodies
    assert "Second duty" in bodies


def test_docx_resolves_tracked_changes_to_the_result_text(tmp_path):
    path = tmp_path / "tracked.docx"
    document = Document()
    paragraph = document.add_paragraph()
    ins = OxmlElement("w:ins")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Accepted target is 2040"
    run.append(text)
    ins.append(run)
    paragraph._p.append(ins)
    delete = OxmlElement("w:del")
    del_run = OxmlElement("w:r")
    del_text = OxmlElement("w:delText")
    del_text.text = "old 2030 wording"
    del_run.append(del_text)
    delete.append(del_run)
    paragraph._p.append(delete)
    document.save(path)
    _mime, blocks = parse_file(path)
    blob = " ".join(block.text for block in blocks)
    assert "Accepted target is 2040" in blob
    assert "2030" not in blob


def test_docx_reads_footnotes(tmp_path):
    path = tmp_path / "notes.docx"
    _write_docx_with_footnote(path, "See the note.", "Scope covers all campuses.")
    _mime, blocks = parse_file(path)
    blob = " ".join(block.text for block in blocks)
    assert "Scope covers all campuses." in blob


def test_docx_reads_text_boxes(tmp_path):
    path = tmp_path / "box.docx"
    document = Document()
    document.add_paragraph("Body text.")
    paragraph = document.add_paragraph()
    txbx = OxmlElement("w:txbxContent")
    inner = OxmlElement("w:p")
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = "Text box: net zero by 2040"
    run.append(node)
    inner.append(run)
    txbx.append(inner)
    paragraph._p.append(txbx)
    document.save(path)
    _mime, blocks = parse_file(path)
    blob = " ".join(block.text for block in blocks)
    assert "net zero by 2040" in blob
    assert "Body text." in blob


def _write_docx_with_footnote(path: Path, body: str, note: str) -> None:
    """Minimal OOXML package with a footnotes part python-docx can open."""
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>{body}</w:t></w:r></w:p>
    <w:sectPr/>
  </w:body>
</w:document>"""
    footnotes_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>
  <w:footnote w:type="continuationSeparator" w:id="0"><w:p/></w:footnote>
  <w:footnote w:id="1"><w:p><w:r><w:t>{note}</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    doc_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>
</Relationships>"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", document_xml)
        zf.writestr("word/_rels/document.xml.rels", doc_rels)
        zf.writestr("word/footnotes.xml", footnotes_xml)
    path.write_bytes(buf.getvalue())
