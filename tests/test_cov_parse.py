"""Branch coverage for rag.parse."""

import sys
import types
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from lxml import etree

from rag.models import IngestError
from rag.parse import (
    _continues_paragraph,
    _docling_extract,
    _docx_text,
    _footnote_texts,
    _markdown_table,
    _outside_tables,
    _pdf_table,
    is_section_title,
    parse_csv,
    parse_docx,
    parse_file,
    parse_md,
    parse_pdf,
    parse_pdf_pages,
    pdf_prose_blocks,
)


def test_parse_file_rejects_unknown_and_unreadable_paths(tmp_path):
    bad = tmp_path / "notes.bin"
    bad.write_text("hello", encoding="utf-8")
    with pytest.raises(IngestError, match="unsupported"):
        parse_file(bad)
    folder = tmp_path / "notes.txt"
    folder.mkdir()
    with pytest.raises(IngestError, match="unreadable"):
        parse_file(folder)


def test_markdown_fences_headings_tables_lists_and_blank_bodies():
    text = (
        "```\n```\n#   \n# Real Title\n|   |   |\n| --- | --- |\n\n"
        "| A | B |\n| --- | --- |\n| 1 | 2 |\n-    \n- item one\n* star\n"
        "1. numbered\n\x00\n\nHello\nworld continues\n\n```\nstill open"
    )
    blocks = parse_md(text, "doc.md")
    kinds = {block.kind for block in blocks}
    assert "code" in kinds
    assert "table" in kinds
    assert "list" in kinds
    assert any("Hello" in block.text and "world continues" in block.text for block in blocks)
    assert _markdown_table(["|  |  |", "| --- | --- |"]) == ""
    assert _continues_paragraph(["hello"], 0) is True
    assert _continues_paragraph(["## Head"], 0) is False
    assert _continues_paragraph(["```"], 0) is False
    assert _continues_paragraph(["- item"], 0) is False
    assert _continues_paragraph(["", "x"], 0) is False
    assert _continues_paragraph(["| a | b |", "| --- | --- |"], 0) is False


def test_csv_header_only_empty_cells_and_extra_columns():
    assert parse_csv("", "empty.csv") == []
    header_only = parse_csv("only\n", "h.csv")
    assert header_only and "only" in header_only[0].text
    assert parse_csv("h1,h2\n,", "blank.csv") == []
    extra = parse_csv("h1\nv,extra", "wide.csv")
    assert "extra" in extra[0].text
    mixed = parse_csv("h1,h2\nv,", "mix.csv")
    assert mixed[0].text == "h1: v"


def test_section_titles_and_pdf_prose_blank_lines():
    assert is_section_title("1. Introduction") is False
    assert is_section_title("Hello") is False
    assert is_section_title("Hello ''") is False
    assert is_section_title("Hello world") is False
    assert is_section_title("Section Title") is True
    blocks = pdf_prose_blocks("Section Title\n\nBody text here", 3)
    assert blocks
    assert any(block.page == 3 for block in blocks)
    assert _pdf_table([]) == ""
    assert _pdf_table([["", None], ["a", "b"]]) == "a | b"


def test_outside_tables_keeps_objects_missing_coordinates():
    keep = _outside_tables([(0.0, 0.0, 10.0, 10.0)])
    assert keep({"x0": None, "x1": 1, "top": 1, "bottom": 1}) is True
    assert keep({"x0": 1, "x1": 2, "top": 1, "bottom": 2}) is False
    assert keep({"x0": 20, "x1": 30, "top": 20, "bottom": 30}) is True


class _Page:
    def __init__(self, text="", boom=False, width=1000, height=1000):
        self._text = text
        self.boom = boom
        self.width = width
        self.height = height
        self.images = []

    def find_tables(self, table_settings=None):
        return []

    def extract_text(self):
        if self.boom:
            raise RuntimeError("broken page")
        return self._text


class _Pdf:
    def __init__(self, pages):
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False


def test_pdf_pages_skip_ocr_when_no_image_and_record_errors(monkeypatch, tmp_path):
    monkeypatch.setitem(sys.modules, "docling", types.ModuleType("docling"))
    _docling_extract(object())
    import pdfplumber

    pages = [_Page(boom=True), _Page(text=""), _Page(text="Visible carbon policy text.")]
    monkeypatch.setattr(pdfplumber, "open", lambda path: _Pdf(pages))
    path = tmp_path / "scan.pdf"
    path.write_bytes(b"%PDF-1.4")
    blocks = parse_pdf(path)
    assert any(block.text.startswith("[page error]") for block in blocks)
    assert any("Visible carbon policy" in block.text for block in blocks)

    direct, errors = parse_pdf_pages([_Page(text="")])
    assert errors == []
    assert direct == []
    supplied, _errors = parse_pdf_pages(
        [_Page(text="Visible carbon policy text.")],
        layout_extract=lambda _page: [],
    )
    assert any("Visible carbon policy" in block.text for block in supplied)


def test_pdf_import_and_open_failures(monkeypatch, tmp_path):
    import pdfplumber

    path = tmp_path / "bad.pdf"
    path.write_bytes(b"%PDF")
    monkeypatch.setitem(sys.modules, "pdfplumber", None)
    with pytest.raises(IngestError, match="pdfplumber"):
        parse_pdf(path)
    monkeypatch.setitem(sys.modules, "pdfplumber", pdfplumber)
    monkeypatch.setattr(pdfplumber, "open", lambda _path: (_ for _ in ()).throw(RuntimeError("corrupt")))
    with pytest.raises(IngestError, match="unreadable"):
        parse_pdf(path)

    def raise_ingest(_path):
        raise IngestError(str(path), "no text")

    monkeypatch.setattr(pdfplumber, "open", raise_ingest)
    with pytest.raises(IngestError, match="no text"):
        parse_pdf(path)


def test_docx_empty_heading_table_tab_deletion_and_textbox(tmp_path):
    path = tmp_path / "edges.docx"
    document = Document()
    style = document.styles.add_style("HeadingX", WD_STYLE_TYPE.PARAGRAPH)
    empty_heading = document.add_paragraph("")
    empty_heading.style = style
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Keep")
    run._r.append(OxmlElement("w:tab"))
    paragraph.add_run("this")
    deleted = OxmlElement("w:del")
    deleted_run = OxmlElement("w:r")
    hidden = OxmlElement("w:t")
    hidden.text = "gone"
    deleted_run.append(hidden)
    deleted.append(deleted_run)
    paragraph._p.append(deleted)
    for box_text in ("   ", "\u00ad"):
        box = OxmlElement("w:txbxContent")
        inner = OxmlElement("w:p")
        inner_run = OxmlElement("w:r")
        node = OxmlElement("w:t")
        node.text = box_text
        inner_run.append(node)
        inner.append(inner_run)
        box.append(inner)
        paragraph._p.append(box)
    document.add_table(rows=1, cols=2)
    document.add_paragraph("Visible body.")
    document.save(path)
    blocks = parse_docx(path)
    blob = " ".join(block.text for block in blocks)
    assert "Keep this" in blob
    assert "gone" not in blob
    assert "Visible body." in blob


def test_docx_skips_a_blank_footnote(monkeypatch, tmp_path):
    monkeypatch.setattr("rag.parse._footnote_texts", lambda _document: ["   ", "Kept note"])
    path = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("Body text.")
    document.save(path)
    blocks = parse_docx(path)
    blob = " ".join(block.text for block in blocks)
    assert "Kept note" in blob
    assert "Body text." in blob


def test_docx_import_and_unreadable(monkeypatch, tmp_path):
    import docx

    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a docx")
    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(IngestError, match="python-docx"):
        parse_docx(path)
    monkeypatch.setitem(sys.modules, "docx", docx)
    monkeypatch.setattr(docx, "Document", lambda _path: (_ for _ in ()).throw(ValueError("corrupt")))
    with pytest.raises(IngestError, match="unreadable"):
        parse_docx(path)


def test_empty_footnote_file_and_blob_parts(tmp_path):
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("Body text.")
    document.save(path)
    blocks = parse_docx(path)
    assert any("Body text." in block.text for block in blocks)

    class Part:
        def __init__(self, name, element=None, blob=None):
            self.partname = name
            self.element = element
            self.blob = blob

    class Package:
        def __init__(self, parts):
            self.parts = parts

    class Doc:
        def __init__(self, parts):
            self.part = types.SimpleNamespace(package=Package(parts))

    xml = """<?xml version="1.0" encoding="UTF-8"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:footnote w:id="-1"/>
  <w:footnote w:id="0"/>
  <w:footnote w:id="1"><w:p><w:r><w:t>Note body</w:t></w:r></w:p></w:footnote>
  <w:footnote w:id="2"><w:p><w:r><w:t>   </w:t></w:r></w:p></w:footnote>
</w:footnotes>"""
    root = etree.fromstring(xml.encode())
    notes = _footnote_texts(
        Doc(
            [
                Part(None),
                Part("/word/document.xml", element=object()),
                Part("/word/footnotes.xml", element=None, blob=None),
                Part("/word/footnotes.xml", element=None, blob=b""),
                Part("/word/footnotes.xml", element=None, blob=xml.encode()),
                Part("/word/footnotes.xml", element=root, blob=None),
            ]
        )
    )
    assert any("Note body" in note for note in notes)

    class Bad:
        @property
        def part(self):
            raise RuntimeError("no package")

    assert _footnote_texts(Bad()) == []
    assert _docx_text(etree.fromstring(xml.encode()))
    assert Path(path).suffix == ".docx"


def test_supplied_layout_and_blank_footnotes(monkeypatch, tmp_path):
    blocks, errors = parse_pdf_pages([_Page(text="Visible policy text here.")], layout_extract=lambda _page: [])
    assert errors == []
    assert any("Visible policy" in block.text for block in blocks)

    monkeypatch.setattr("rag.parse._footnote_texts", lambda _document: ["   ", "A real footnote"])
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("Body sentence.")
    document.save(path)
    parsed = parse_docx(path)
    assert any("A real footnote" in block.text for block in parsed)
    assert all(block.text.strip() for block in parsed)
